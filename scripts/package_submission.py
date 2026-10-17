"""Populate the original template with measured artifacts and honest disclosures."""
from datetime import datetime, timezone
import hashlib
import html
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import textwrap
import zipfile

from docx import Document
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / 'artifacts'
SUB = ROOT / 'submission'

def read(name):
    return json.loads((ROOT/name).read_text(encoding='utf-8'))

def fmt(value, places=3):
    return f'{value:.{places}f}' if value is not None else 'N/A — no valid observations'

def pct(value):
    return f'{100*value:.1f}%' if value is not None else 'N/A — no valid observations'

def pretty(value):
    return json.dumps(value, indent=2, ensure_ascii=False)

def draw_diagram(path, kind):
    image = Image.new('RGB', (1400, 850), '#f2f5fa')
    draw = ImageDraw.Draw(image)
    fontpath = Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts/segoeui.ttf'
    font = ImageFont.truetype(str(fontpath), 23) if fontpath.exists() else ImageFont.load_default()
    titlefont = ImageFont.truetype(str(fontpath), 33) if fontpath.exists() else font
    def box(x,y,w,h,text,color='#ffffff'):
        draw.rounded_rectangle((x,y,x+w,y+h), 14, fill=color, outline='#28496b', width=2)
        lines = []
        for paragraph in text.split('\n'):
            lines.extend(textwrap.wrap(paragraph, max(10,int(w/12))) or [''])
        for i,line in enumerate(lines):
            draw.text((x+w/2,y+h/2+(i-(len(lines)-1)/2)*28),line,font=font,fill='#14324c',anchor='mm')
    def arrow(points):
        draw.line(points, fill='#456c93', width=4)
        x,y=points[-1]; px,py=points[-2]
        import math
        angle=math.atan2(y-py,x-px)
        draw.polygon([(x,y),(x-15*math.cos(angle-.5),y-15*math.sin(angle-.5)),(x-15*math.cos(angle+.5),y-15*math.sin(angle+.5))],fill='#456c93')
    if kind == 'rag':
        draw.text((45,25),'Problem 1 | Local, cited question answering',font=titlefont,fill='#14324c')
        left=['PDF / HTML / Markdown','Extract text + PDF page','Chunk: 120 words / overlap 24','MiniLM-L6-v2 / 384 float32','sqlite-vec + source, page, category']
        right=['HTTP / CLI question','Embed query; category prefilter','Top-k cosine search (default 3)','Similarity threshold: 0.38','Qwen answer + source indices','Validate citations; answer + usage']
        for i,text in enumerate(left):
            y=105+i*134; box(45,y,490,94,text)
            if i<4: arrow([(290,y+94),(290,y+134)])
        for i,text in enumerate(right):
            y=105+i*114; box(620,y,440,80,text,'#e6f2ff')
            if i<5: arrow([(840,y+80),(840,y+114)])
        arrow([(535,690),(575,690),(575,373),(620,373)])
        box(1100,390,260,195,'No relevant context\nor invalid citations\nAbstain\n0 generation tokens when no context','#fff0d7')
        arrow([(1060,487),(1100,487)])
        arrow([(1060,715),(1230,715),(1230,585)])
    else:
        draw.text((45,25),'Problem 2 | Auditable, independent LLM judging',font=titlefont,fill='#14324c')
        left=['JSON / YAML suite + references','Validate cases; hide model identity','Rubric + labeled evidence + anchors','SmolLM2 (different from generator)','Raw prompts / responses / token log']
        right=['Pointwise score OR A/B both orders','Parse JSON + schema validation','Malformed: safe repair + one retry','Invalid again: explicit error','Remap A/B; disagreement = abstain','Aggregate + bias / gold / A/B reports']
        for i,text in enumerate(left):
            y=105+i*134; box(45,y,535,94,text)
            if i<4: arrow([(312,y+94),(312,y+134)])
        for i,text in enumerate(right):
            y=105+i*114; box(760,y,590,80,text,'#e6f2ff')
            if i<5: arrow([(1055,y+80),(1055,y+114)])
        arrow([(580,420),(670,420),(670,145),(760,145)])
        arrow([(760,259),(680,259),(680,690),(580,690)])
        draw.text((50,805),'Scores 0–4 | explicit rationale per criterion | consensus coverage always reported',font=font,fill='#456c93')
    image.save(path)

def build_evidence(rag, ab, validation, idempotency):
    query = next(c['result'] for c in rag['cases'] if not c['result']['abstained'])
    queryview = {key: query[key] for key in ['query_id','question','answer','chunk_count','citations','usage','retrieval_ms','generation_ms','end_to_end_ms'] if key in query}
    records = [json.loads(line) for line in (ART/'llm_audit.jsonl').read_text(encoding='utf-8').splitlines()]
    request_id = validation['cases'][0]['baseline']['request_ids'][0]
    record = next(r for r in records if r['request_id'] == request_id)
    auditview = {'request_id':record['request_id'], 'tag':record['tag'], 'model':record['model'],
                 'messages':record['request']['messages'], 'raw_response':json.loads(record['raw_response']), 'latency_ms':record['latency_ms']}
    case = ab['cases'][0]
    pair=case['mitigated']
    compact=lambda verdict:{key:verdict[key] for key in ['winner','canonical_winner','rationale','request_ids'] if key in verdict}
    positionview = {'case_id':case['id'], 'forward':compact(pair['forward']), 'reverse':compact(pair['reverse']),
                    'flipped':pair['flipped'], 'consensus_winner':pair['winner'], 'suite_position_bias':ab['position_bias']}
    adversarial = next(c for c in validation['cases'] if c['id'] == 'verbose-wrong')
    cards=[('query','End-to-end RAG query',queryview,'artifacts/rag_results.json'),
           ('idempotency','Idempotent re-ingestion',idempotency,'artifacts/idempotency.json'),
           ('judge','Raw judge prompt + response',auditview,'artifacts/llm_audit.jsonl'),
           ('position','Same pair in both orders',positionview,'artifacts/ab_comparison.json'),
           ('adversarial','Verbose-but-wrong adversarial verdict',adversarial,'artifacts/judge_validation.json')]
    content='<!doctype html><html><meta charset="utf-8"><title>Measured assignment evidence</title><style>body{background:#edf1f7;margin:24px;font-family:Segoe UI,Arial;color:#122e46}section{background:white;border:1px solid #cbd5e1;border-radius:12px;padding:28px;margin:24px 0;width:1120px;box-sizing:border-box}h1{font-size:24px;margin:0 0 8px}.meta{font-size:13px;color:#52677c;margin-bottom:16px}pre{font:13px/1.45 Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;margin:0;padding:18px;background:#f5f8fc;border-radius:8px}</style>'
    for id,title,data,source in cards:
        content+=f'<section id="{id}"><h1>{html.escape(title)}</h1><div class="meta">Screenshot of measured records · source: {html.escape(source)} · values shown without alteration</div><pre>{html.escape(pretty(data))}</pre></section>'
    content+='</html>'
    (ART/'evidence.html').write_text(content,encoding='utf-8')
    result=subprocess.run([sys.executable,str(ROOT/'scripts/capture_evidence.py')],cwd=ROOT,text=True,capture_output=True)
    (ART/'screenshot_capture.txt').write_text(result.stdout+result.stderr,encoding='utf-8')
    if result.returncode:
        if not all((ART/(id+'.png')).exists() for id,*_ in cards):
            raise RuntimeError('Evidence screenshot capture failed: '+result.stderr)
        print('Retaining previously captured evidence screenshots; see screenshot_capture.txt.')

def fill_table(table, rows):
    while len(table.rows)<len(rows): table.add_row()
    for row,values in zip(table.rows,rows):
        for cell,text in zip(row.cells,values):
            cell.text=str(text)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4)
                for run in paragraph.runs: run.font.size=Pt(9)

def image_cell(table, path, caption):
    cell=table.cell(0,0)
    cell.text=caption
    image=Image.open(path)
    w,h=image.size
    width=min(6.0,6.5*w/h)
    cell.add_paragraph().add_run().add_picture(str(path),width=Inches(width))

def make_document(rag, ab, val, costs, ablation):
    source=SUB/'source/submission_template.docx'
    candidate_overrides={}
    existing=SUB/'completed_submission.docx'
    if existing.exists():
        prior=Document(existing)
        for index,element in enumerate(list(prior.element.body)):
            if index in (77,81,152,156) and element.tag.endswith('}tbl'):
                text=Table(element,prior).cell(0,0).text
                if text.strip() and 'Candidate-written reflection required' not in text and 'Candidate interpretation required' not in text:
                    candidate_overrides[index]=text
    document=Document(source)
    elements=list(document.element.body)
    tables={i:Table(element,document) for i,element in enumerate(elements) if element.tag.endswith('}tbl')}
    paragraphs={i:Paragraph(element,document) for i,element in enumerate(elements) if element.tag.endswith('}p')}
    summary=rag['summary']; retrieval=summary['retrieval']; quality=summary['answer_quality']; latency=summary['latency']
    env=rag['environment']
    variables='RAG_DB, EMBEDDING_MODEL, CHUNK_SIZE, CHUNK_OVERLAP, MIN_SIMILARITY, TOP_K, GENERATOR_BASE_URL, GENERATOR_MODEL, GENERATOR_FAMILY, GENERATOR_API_KEY, JUDGE_BASE_URL, JUDGE_MODEL, JUDGE_FAMILY, JUDGE_API_KEY, LLM_TIMEOUT, MODEL_THREADS, AUDIT_LOG, QUERY_LOG. Defaults in .env.example; no keys needed locally.'
    prerequisites=f'Python 3.13; Windows x64 automatic setup; 4 GB free RAM and 5 GB disk recommended; initial Internet download. Tested Python {env["python"]}, {env["ram_gib"]} GiB total RAM, {env["logical_cpus"]} logical CPUs; CPU inference, four threads/server. Linux/macOS can use independently configured local llama.cpp endpoints.'
    install='From the extracted project folder (local repository included; no public URL was supplied):\npython scripts/bootstrap.py\n.\\.venv\\Scripts\\python.exe main.py models start\nInitial models are about 1.5 GB; download time depends on the connection. Cached setup is fast; full evaluation takes longer than setup.'
    disclosure='Codex (OpenAI coding assistant) generated and revised the implementation, synthetic fixtures, technical documentation, tests, evaluation harness and submission artifacts. Real local Qwen/SmolLM model calls produced the measured results. No candidate-authored implementation or personal reflection is claimed. The candidate must independently review the work and write the required reflections.'
    personal='Candidate-written reflection required. The template explicitly requires this answer without AI assistance. No personal account has been fabricated; the candidate must write this section before submitting.'
    fields=[['Field','Your entry'],['Full name','Abhishek — workspace name only; full candidate name was not supplied.'],
        ['Email','Not supplied in the assignment materials; candidate must enter their contact email.'],
        ['Date submitted',datetime.now().strftime('%Y-%m-%d')+' — artifact preparation date; actual submission date is candidate-controlled.'],
        ['Repository link','Local Git repository in this folder; submission/source and project archive included. No remote repository published.'],
        ['Total time spent','AI-assisted execution; machine runtime and artifact timestamps recorded. Candidate personal time was not supplied.'],
        ['Problem 1 — vector store chosen','sqlite-vec 0.1.6'],
        ['Problem 2 — judge model / generator model','Judge: SmolLM2-1.7B-Instruct Q4_K_M (HuggingFaceTB); generator: Qwen2.5-0.5B-Instruct Q4_K_M (Qwen).']]
    fill_table(tables[4],fields)
    image_cell(tables[16],ART/'rag_architecture.png','Full RAG path, category prefilter, threshold and no-context/citation-failure branches.')
    for index,prefix,value in [(20,'Prerequisites',prerequisites),(21,'Environment variables',variables),
        (31,'Vector store chosen + one-line why','sqlite-vec: local, transactional, no separate server, native metadata-prefiltered cosine search.'),
        (36,'k used','k = 3; fixed 20 questions (16 answerable + 4 negatives).'),
        (96,'Prerequisites',prerequisites),(97,'Environment variables',variables),
        (107,'Judge & generator configured independently?','Yes. Separate GENERATOR_* and JUDGE_* URL, model, family and optional key settings. Same declared family is rejected.'),
        (111,'Judging mode used','Pointwise reference-based scoring plus pairwise A/B in both orders. Missing references are supported but do not establish factual truth.'),
        (125,'Declared winner + one-line justification',f'{ab["winner"]}: higher consensus pairwise wins across all 8 generated pairs; disagreements abstain and ties earn neither a win. A={pct(ab["A"]["win_rate_all_pairs"])}, B={pct(ab["B"]["win_rate_all_pairs"])}.')]:
        paragraphs[index].text=prefix+'  '+value
    fill_table(tables[23],[[install]])
    fill_table(tables[26],[['.\\.venv\\Scripts\\python.exe main.py ingest corpus\n# PDF/HTML/MD; defaults: 120 words, overlap 24.\n# Change both explicitly: main.py ingest corpus --chunk-size 60 --overlap 12']])
    fill_table(tables[29],[['.\\.venv\\Scripts\\python.exe main.py query "How long do I have to request a refund?" --category billing\n.\\.venv\\Scripts\\python.exe main.py serve --port 8080\n# In another PowerShell terminal:\nInvoke-RestMethod -Uri http://127.0.0.1:8080/query -Method Post -ContentType \'application/json\' -Body \'{"question":"Which encryption protects data at rest?","k":3,"category":"security"}\'']])
    fill_table(tables[37],[['Metric','Value','How computed / notes'],
        ['Recall@k / Hit Rate',f'{fmt(retrieval["recall_at_k"])} / {fmt(retrieval["hit_rate"])}','Macro average over 16 answerable cases; source + normalized evidence substring defines relevant chunks.'],
        ['MRR',fmt(retrieval['mrr']),'Inverse rank of first relevant chunk; zero if absent.'],
        ['nDCG@k',fmt(retrieval['ndcg_at_k']),'Binary gains, log2 rank discount; ideal ranking over min(k, relevant count).'],
        ['Context precision',fmt(retrieval['context_precision']),'Relevant returned / number returned. Overlapping relevant chunks both count; negative cases reported separately.']])
    fill_table(tables[40],[['Metric','Score','Method','Notes'],
        ['Faithfulness / groundedness',fmt(quality['faithfulness_mean_0_to_4'])+' / 4','Independent LLM judge','SmolLM2 sees answer, context, gold; inspect rationales. Small-model validity limits apply.'],
        ['Answer relevance',fmt(quality['answer_relevance_mean_0_to_4'])+' / 4','Independent LLM judge',f'Explicit relevance criterion; judge errors: {quality["judge_errors"]}.'],
        ['EM',pct(quality['em']),'Frozen gold answers',f'20 cases; answerable-only={pct(quality["answerable_em"])}. Citation/case/punctuation normalization.'],
        ['F1',fmt(quality['f1']),'Multiset token overlap',f'20 cases; answerable-only={fmt(quality["answerable_f1"])}. Negative abstention={pct(summary["negative_abstention_rate"])}.']])
    costrows=[['Vectors','Your store ($/mo)','Managed DB ($/mo)','Savings / notes']]
    for row,label in zip(costs['rows'],['100K','1M','10M']):
        costrows.append([label,fmt(row['self_host_with_weekly_backup_usd'],2),fmt(row['managed_cheapest_usd'],2),
            f'{row["managed_cheapest_tier"]}; savings ${row["savings_vs_cheapest_usd"]:.2f}. Standard-only ${row["managed_standard_usd"]:.2f}. Storage {row["namespace_gb"]:.3f} GB.'])
    fill_table(tables[44],costrows)
    assumptions='Scenario estimates, not deployed-scale benchmarks. 384 x float32 + 20-byte ID + 768 bytes text/metadata; 1,000 serial queries/month, 1% records updated monthly, k=3. US East: DigitalOcean NYC1 single VM (1/2/4 GiB RAM, 25/50/80 GiB SSD) at $6/$12/$24 + 20% weekly backups. SQLite disk overhead modeled 2x; no HA. Managed: Pinecone AWS us-east-1, eligible Starter/Builder/Standard; $0/$20/$50 minimum respectively, $0.33/GB-month, $16/M RU, $4/M WU frozen scenario inputs. RU/query=max(0.25, namespace GB); updates account for old+new bytes. Initial loading, LLM/embedding compute, labor, tax excluded from recurring table. Provider pricing and quota sources + calculations are in configs/cost_assumptions.json and artifacts/cost_comparison.json. At 100K managed free wins; at 10M exact scans can have poor latency. No latency extrapolation or equal-reliability claim.'
    fill_table(tables[45],[[assumptions]])
    fill_table(tables[48],[['Metric','Value (ms)','Notes / conditions'],
        ['Retrieval p50',fmt(latency['retrieval_p50_ms'],2),f'k=3, {rag["settings"]["chunks"]} chunks; 100 warm serial samples; query embedding + sqlite overhead included.'],
        ['Retrieval p95',fmt(latency['retrieval_p95_ms'],2),f'5 passes x 20 queries; {env["logical_cpus"]} logical CPUs, {env["ram_gib"]} GiB RAM.'],
        ['End-to-end p95 (optional)',fmt(latency['end_to_end_p95_ms'],2),'20 actual RAG calls, including generation; excludes separate evaluation judge.']])
    ablationtext='; '.join(f'{r["size_words"]}/{r["overlap_words"]} words: {r["chunks"]} chunks, MRR={r["retrieval"]["mrr"]:.3f}, precision={r["retrieval"]["context_precision"]:.3f}' for r in ablation)
    design1={53:'sqlite-vec keeps vectors and metadata in one local transaction and avoids a daemon, Docker, or cloud account. FAISS would require separate metadata persistence; Qdrant/pgvector offer stronger scale and concurrency at greater setup cost. Those alternatives were considered, not falsely benchmarked. Exact scan is acceptable for this tiny corpus.',
        55:'Default 120 words / 24 overlap balances duplicated context against splitting support facts. Measured ablation: '+ablationtext+'. These tiny-fixture measurements are not proof of a universal optimum. Words are not tokenizer tokens, so long technical tokens can still truncate.',
        57:'MiniLM-L6-v2 produces 384 float32 values (1,536 raw vector bytes). CPU ONNX inference is free of API charges and much smaller than a GPU stack. English support QA is a limited domain; stronger/multilingual embeddings require a new index and independent evaluation.',
        59:f'Cosine similarity below 0.38 yields deterministic abstention with zero generator tokens. The generator also has an explicit abstention instruction; malformed or invalid citation indices are rejected. Measured negative abstention={pct(summary["negative_abstention_rate"])}; positive abstention={pct(summary["answerable_abstention_rate"])}. Thresholding is a heuristic and cannot prove semantic support.',
        61:'A source identity derives from corpus root plus relative path; a fingerprint includes file bytes, model and chunk settings. Unchanged files skip embedding. Fully prepared replacements delete old chunks/vectors and insert unique chunk hashes in one transaction. Tests verify unchanged IDs/count, replacement and rollback on embedding failure. Re-ingestion does not prune deleted files.',
        63:'Accepted limits: single writer, local failure domain, exact scans, small synthetic benchmark, small generator/judge and no OCR. At 10M the estimated index exceeds selected RAM and exact scans may violate a latency target. Move to ANN self-hosting or managed service for sustained concurrency, latency SLAs, replication/recovery needs, or operations costs beyond savings.'}
    for index,text in design1.items(): fill_table(tables[index],[[text]])
    image_cell(tables[69],ART/'query.png','Actual query log: cited answer, latency, chunk count and model-reported usage.')
    image_cell(tables[73],ART/'idempotency.png','Actual ingestion records; unchanged vector count after a second identical ingestion.')
    fill_table(tables[77],[[personal]])
    fill_table(tables[81],[[personal+' Technical evidence for independent review: early raw audit records show missing generator citations; the implementation now requests source indices in structured JSON. This is an AI-observed implementation note, not a candidate reflection.']])
    fill_table(tables[85],[[disclosure]])
    image_cell(tables[92],ART/'judge_architecture.png','Structured parsing and fallback, independent judge, A/B swap and consensus aggregation.')
    fill_table(tables[99],[[install]])
    fill_table(tables[102],[['.\\.venv\\Scripts\\python.exe main.py judge data/example_suite.yaml --output artifacts/example_report.json\n.\\.venv\\Scripts\\python.exe main.py judge data/judge_probes.json --output artifacts/probe_suite_report.json']])
    fill_table(tables[105],[['.\\.venv\\Scripts\\python.exe main.py evaluate generate-ab\n.\\.venv\\Scripts\\python.exe main.py compare data/comparison_suite.json --output artifacts/ab_comparison.json\n# A: ungrounded-v1; B: grounded-v2; same Qwen generator.\n# Frozen generated outputs can be rejudged without regeneration.']])
    from src.judge import RUBRIC
    rubricrows=[['Criterion','Definition / score anchors','Weight']]
    for name,data in RUBRIC.items():
        rubricrows.append([name.replace('_',' ').capitalize(),data['definition']+' Anchors: 0 unsupported/wrong, 1 major errors, 2 mixed, 3 minor omission, 4 complete/supportable.',fmt(data['weight'],2)])
    fill_table(tables[113],rubricrows)
    position=ab['position_bias']; lengths=val['verbosity_and_style']; identity=val['self_enhancement']
    fill_table(tables[117],[['Bias','Mitigation implemented (in code)','Metric & result (before → after)'],
        ['Position (A/B order)','Judge.pairwise: both orders, canonical remapping, disagreement -> abstain; calibrated vs baseline on same 8 pairs.',f'Raw flip rate {pct(position["baseline"]["flip_rate"])} -> {pct(position["mitigated"]["flip_rate"])}; consensus coverage {pct(position["baseline"]["consensus_coverage"])} -> {pct(position["mitigated"]["consensus_coverage"])}. No zero-bias claim from abstention.'],
        ['Verbosity / length','Anchors + unsupported-padding penalty; same answer with neutral padding.',f'Absolute score drift {fmt(lengths["baseline"]["absolute_padding_score_drift"])} -> {fmt(lengths["mitigated"]["absolute_padding_score_drift"])} (one probe; 0–4 scale).'],
        ['Self-enhancement','Different generator/judge families + omit provenance from model prompt.',f'Label-only score gap {fmt(identity["score_gap_exposed"])} -> {fmt(identity["score_gap_blind"])}. One counterfactual pair, not broad bias validation.'],
        ['Sycophancy / style','Per-criterion evidence rationales; confidently wrong and injection probes.',f'Confident wrong accepted: {lengths["baseline"]["confidently_wrong_passed"]} -> {lengths["mitigated"]["confidently_wrong_passed"]}. Small sample; inspect raw verdicts.'],
        ['Score clustering','Few-shot anchors; pairwise mode for A/B.',f'Score population stddev {fmt(val["baseline"]["score_stddev"])} -> {fmt(val["mitigated"]["score_stddev"])} across same 8 probes; dispersion alone is not validity.']])
    adversarial=next(c for c in val['cases'] if c['id']=='verbose-wrong')
    fooled=adversarial['mitigated']['passed']
    fill_table(tables[121],[['Method','Result','Notes'],
        ['Agreement with human / gold',pct(val['mitigated']['agreement']),f'8 synthetic authored binary gold labels, not a human review study. Before={pct(val["baseline"]["agreement"])}.'],
        ["Cohen’s kappa / correlation",fmt(val['mitigated']['cohens_kappa']),f'Binary pass/fail; baseline kappa={fmt(val["baseline"]["cohens_kappa"])}. Parse errors count as failures.'],
        ['Test-retest flip rate',pct(val['test_retest']['flip_rate']),'4 repeated cases, temperature=0, seed=42. Small-sample consistency does not establish correctness.'],
        ['Adversarial probe outcome','Fooled: '+('YES' if fooled else 'NO'),'Verbose-but-wrong probe; full rationale/score and other probes in judge_validation.json.']])
    fill_table(tables[124],[['Config','Pass rate','Mean score','Win rate'],
        ['Config A — ungrounded-v1',pct(ab['A']['pass_rate']),fmt(ab['A']['mean_score'])+' / 4',pct(ab['A']['win_rate_all_pairs'])],
        ['Config B — grounded-v2',pct(ab['B']['pass_rate']),fmt(ab['B']['mean_score'])+' / 4',pct(ab['B']['win_rate_all_pairs'])]])
    design2={130:'Pointwise reference-based mode provides per-criterion failure explanations for QA with gold facts. Pairwise mode compares two prompt configurations without assuming the 0–4 scale is perfectly calibrated. Reference-free mode is supported for constraints/style but cannot establish factual truth without external evidence.',
        132:'Constrain output by JSON schema, then independently parse and validate exact criteria, integer scores, bounds and nonempty rationales. Fallbacks recover embedded/fenced JSON, trailing commas and safe literal dictionaries. One model repair call is allowed; another invalid verdict becomes an explicit error and remains in the denominator. Prompt, raw output and both calls are logged.',
        134:'Generator is Qwen2.5, judge is SmolLM2 from a different publisher/model family. Independent URL/model/family settings are enforced; provenance is stripped before judging. This reduces direct self-preference and model-label sensitivity but does not eliminate correlated training-data or style biases. The one-pair identity probe is explicitly narrow.',
        136:f'Position sensitivity directly changes A/B outcomes, so both orders are mandatory. Observed raw flips {pct(position["baseline"]["flip_rate"])} -> {pct(position["mitigated"]["flip_rate"])} across only 8 pairs. Consensus abstention prevents an order-disagreeing comparison from choosing a winner; coverage {pct(position["mitigated"]["consensus_coverage"])} must be considered. This evidence is too small to claim bias elimination.',
        138:f'No autonomous production release gate on this evidence. The small local judge achieves {pct(val["mitigated"]["agreement"])} agreement on 8 synthetic labels and may produce inconsistent rationales. Use it as a review aid; require human review for disagreements, high-impact decisions, adversarial failures, and score changes near the pass threshold. Validate on an independently labeled production set before expanding its role.'}
    for index,text in design2.items(): fill_table(tables[index],[[text]])
    image_cell(tables[144],ART/'judge.png','Actual audited judge prompt, exact raw response and token usage.')
    image_cell(tables[148],ART/'position.png','Actual forward/reverse comparison, canonical winners, flip rates and coverage.')
    fill_table(tables[152],[[personal]])
    fill_table(tables[156],[[pretty(adversarial)+'\n\nCandidate interpretation required without AI assistance. The recorded verdict above is actual pipeline output; it is not a personal reflection.']])
    fill_table(tables[160],[[disclosure]])
    for index,text in candidate_overrides.items():
        fill_table(tables[index],[[text]])
    # Let the filled tables grow with their content instead of inheriting empty-box heights.
    for table in document.tables:
        for row in table.rows:
            for height in list(row._tr.xpath('./w:trPr/w:trHeight')):
                height.getparent().remove(height)
    output=SUB/'completed_submission.docx'
    document.save(output)
    original=Document(source)
    expected=[p.text for p in original.paragraphs if re.match(r'^[12]\.\d',p.text)]
    produced=[p.text for p in Document(output).paragraphs if re.match(r'^[12]\.\d',p.text)]
    if expected!=produced: raise RuntimeError('Original section headers changed')
    texts=[p.text for p in document.paragraphs]+[c.text for t in document.tables for row in t.rows for c in row.cells]
    if any(re.search(r'\[\s*(?:\.\.\.|your answer|paste|replace|name)',text,re.I) for text in texts):
        raise RuntimeError('Unreplaced template placeholder found')
    (ART/'template_validation.json').write_text(pretty({'section_headers_preserved':True,'section_header_count':len(expected),
        'unreplaced_placeholders':0,'personal_reflections':'Candidate-required, explicitly marked; no personal authorship claimed.'}),encoding='utf-8')

def results_markdown(rag, ab, val, costs, ablation):
    r=rag['summary']; q=r['answer_quality']; l=r['latency']; p=ab['position_bias']
    text=f'''# Measured results

Generated from committed JSON artifacts. Synthetic dataset: 8 support documents,
20 QA cases (16 answerable + 4 negatives); 8 generated A/B pairs and 8 authored
judge-validation labels. These are development fixtures, not a production or human study.

## Problem 1

| Metric | Measured value |
|---|---:|
| Recall@3 | {fmt(r['retrieval']['recall_at_k'])} |
| Hit rate@3 | {fmt(r['retrieval']['hit_rate'])} |
| MRR | {fmt(r['retrieval']['mrr'])} |
| nDCG@3 | {fmt(r['retrieval']['ndcg_at_k'])} |
| Context precision | {fmt(r['retrieval']['context_precision'])} |
| Faithfulness (local judge, 0–4) | {fmt(q['faithfulness_mean_0_to_4'])} |
| Answer relevance (local judge, 0–4) | {fmt(q['answer_relevance_mean_0_to_4'])} |
| Exact match (all 20) | {pct(q['em'])} |
| Token F1 (all 20) | {fmt(q['f1'])} |
| Answerable-only EM / F1 | {pct(q['answerable_em'])} / {fmt(q['answerable_f1'])} |
| Negative abstention rate | {pct(r['negative_abstention_rate'])} |
| Positive abstention rate | {pct(r['answerable_abstention_rate'])} |
| Retrieval p50 / p95 | {fmt(l['retrieval_p50_ms'],2)} / {fmt(l['retrieval_p95_ms'],2)} ms |
| End-to-end p95 | {fmt(l['end_to_end_p95_ms'],2)} ms |

Retrieval includes embedding and SQLite connection/search overhead. 100 warm serial
samples; end-to-end uses 20 actual calls. The corpus has {rag['settings']['chunks']}
chunks. Hardware and runtime are recorded in `rag_results.json`.
Judge errors: {q['judge_errors']}. LLM quality scores are small-model judgments,
not independent human verification. Citation indices exist, but that alone does
not prove every cited passage semantically supports every generated claim.

Retrieval vs generation: retrieval hit rate is {pct(r['retrieval']['hit_rate'])},
while answerable exact match is {pct(q['answerable_em'])}. The raw per-case results
show whether the remaining misses arise after a relevant chunk was found.
Short-answer EM is strict about wording; F1 gives a complementary lexical measure.
Neither metric proves multi-hop or out-of-domain quality.

### Chunking ablation

| Size / overlap (words) | Chunks | MRR | Context precision |
|---|---:|---:|---:|
'''
    for row in ablation:
        text+=f'| {row["size_words"]} / {row["overlap_words"]} | {row["chunks"]} | {fmt(row["retrieval"]["mrr"])} | {fmt(row["retrieval"]["context_precision"])} |\n'
    text+='''
### Store cost projections (USD/month)

| Vectors | sqlite-vec VM + weekly backup | Cheapest eligible managed | Savings | Managed Standard |
|---|---:|---:|---:|---:|
'''
    for row in costs['rows']:
        text+=f'| {row["vectors"]:,} | ${row["self_host_with_weekly_backup_usd"]:.2f} | ${row["managed_cheapest_usd"]:.2f} ({row["managed_cheapest_tier"]}) | ${row["savings_vs_cheapest_usd"]:.2f} | ${row["managed_standard_usd"]:.2f} |\n'
    text+='''
All assumptions and arithmetic are in `cost_comparison.json`. Initial managed
loading is shown separately there. Rates are frozen scenario inputs based on
[Pinecone pricing](https://www.pinecone.io/pricing/),
[metering rules](https://docs.pinecone.io/guides/manage-cost/understanding-cost), and
[DigitalOcean VM/backup pricing](https://www.digitalocean.com/pricing/droplets).
At small scale a free managed tier is cheaper. At 10M exact cosine scans can be
too slow despite fitting on disk. These estimates do not include LLM hardware,
staff labor, taxes or equal availability. No 10M-vector latency claim is made.

## Problem 2

| Config | Pass rate | Mean score (0–4) | Consensus win rate, all pairs |
|---|---:|---:|---:|
'''
    for label,name in [('A','Ungrounded-v1'),('B','Grounded-v2')]:
        text+=f'| {label}: {name} | {pct(ab[label]["pass_rate"])} | {fmt(ab[label]["mean_score"])} | {pct(ab[label]["win_rate_all_pairs"])} |\n'
    text+=f'''
Declared winner: **{ab['winner']}**, using higher two-order consensus win count.
Ties/disagreements earn neither config a win; the rule never forces a winner.

| Bias / validation metric | Before | After |
|---|---:|---:|
| Position order-disagreement rate (8 pairs) | {pct(p['baseline']['flip_rate'])} | {pct(p['mitigated']['flip_rate'])} |
| Position consensus coverage | {pct(p['baseline']['consensus_coverage'])} | {pct(p['mitigated']['consensus_coverage'])} |
| Padding absolute score drift (1 probe) | {fmt(val['verbosity_and_style']['baseline']['absolute_padding_score_drift'])} | {fmt(val['verbosity_and_style']['mitigated']['absolute_padding_score_drift'])} |
| Identity label score gap (1 probe pair) | {fmt(val['self_enhancement']['score_gap_exposed'])} | {fmt(val['self_enhancement']['score_gap_blind'])} |
| Confidently wrong answer accepted | {val['verbosity_and_style']['baseline']['confidently_wrong_passed']} | {val['verbosity_and_style']['mitigated']['confidently_wrong_passed']} |
| Score population standard deviation | {fmt(val['baseline']['score_stddev'])} | {fmt(val['mitigated']['score_stddev'])} |
| Gold agreement (8 synthetic labels) | {pct(val['baseline']['agreement'])} | {pct(val['mitigated']['agreement'])} |
| Cohen's kappa | {fmt(val['baseline']['cohens_kappa'])} | {fmt(val['mitigated']['cohens_kappa'])} |

Test-retest: {pct(val['test_retest']['flip_rate'])} pass/fail flips over
{val['test_retest']['n']} cases at temperature 0, seed 42.
Before/after is the same fixed probe set with an uncalibrated versus anchored
prompt, except the identity probe which compares exposed versus hidden labels.
Some mitigations may fail or worsen a metric; the table reports observed values.
Consensus abstention controls unreliable decisions; it does not establish that
the underlying model is unbiased. Small sample sizes limit conclusions.

The judge is suitable as a review aid, not an autonomous production release gate.
Human review is required for disagreements, high-impact outputs and adversarial
failures. Wider score spread is not sufficient evidence of better calibration.

## Audit and provenance

All judge/generator prompts, raw model responses and real token counts are in
`llm_audit.jsonl`; request IDs link each report to its calls. Earlier debugging
runs remain in the log and are not pooled into final metrics. To inspect or
rerun one call, use `main.py replay REQUEST_ID` or add `--rerun`.
`model_manifest.json` records exact model revisions and SHA256 hashes.
`template_validation.json`, `test_results.txt` and `http_smoke.json` record checks.
The evidence screenshots show actual audit records in the local evidence viewer.

Candidate identity/contact and mandatory independent reflections remain for the
candidate. Codex generated the implementation and technical write-up; no personal
authorship or human validation is claimed.
'''
    (ART/'RESULTS.md').write_text(text,encoding='utf-8')

def main():
    os.chdir(ROOT)
    SUB.mkdir(exist_ok=True); (SUB/'source').mkdir(exist_ok=True)
    for filename in ['Gen AI_assignment (1).pdf','submission_template.docx']:
        target=SUB/'source'/filename
        if not target.exists():
            shutil.copy2(Path.home()/'Downloads'/filename,target)
    from scripts.costs import calculate
    costs=calculate()
    rag=read('artifacts/rag_results.json')
    ab=read('artifacts/ab_comparison.json')
    val=read('artifacts/judge_validation.json')
    ablation=read('artifacts/chunk_ablation.json')
    if 'summary' not in rag or 'winner' not in ab or 'test_retest' not in val:
        raise RuntimeError('Complete main.py evaluate all before packaging')
    if (ROOT/'.runtime/manifest.json').exists():
        shutil.copy2(ROOT/'.runtime/manifest.json',ART/'model_manifest.json')
        shutil.copy2(ROOT/'.runtime/manifest.json',ROOT/'configs/runtime.lock.json')
    freeze=subprocess.run([sys.executable,'-m','pip','freeze'],capture_output=True,text=True,check=True)
    (ROOT/'requirements.lock.txt').write_text(freeze.stdout,encoding='utf-8')
    draw_diagram(ART/'rag_architecture.png','rag')
    draw_diagram(ART/'judge_architecture.png','judge')
    build_evidence(rag,ab,val,read('artifacts/idempotency.json'))
    results_markdown(rag,ab,val,costs,ablation)
    make_document(rag,ab,val,costs,ablation)
    required={
        'Problem 1 ingestion / metadata / citations':'src/ingestion.py, src/store.py, src/rag.py; artifacts/http_smoke.json',
        'Problem 1 retrieval / quality / latency':'artifacts/rag_results.json; data/rag_questions.json',
        'Problem 1 scale costs':'configs/cost_assumptions.json; artifacts/cost_comparison.json',
        'Problem 1 idempotency':'artifacts/idempotency.json; artifacts/idempotency.png',
        'Problem 2 rubric / structured fallback':'src/judge.py; artifacts/test_results.txt',
        'Problem 2 bias / validation':'artifacts/judge_validation.json; artifacts/ab_comparison.json',
        'Problem 2 generated A/B outputs':'data/comparison_suite.json; artifacts/generation_usage.json',
        'Problem 2 exact prompt / raw response / usage':'artifacts/llm_audit.jsonl; main.py replay REQUEST_ID',
        'Architecture / screenshots':'docs/*.mmd; artifacts/*_architecture.png; artifacts/evidence.html; evidence PNGs',
        'Original template populated':'submission/completed_submission.docx; artifacts/template_validation.json',
        'Candidate-only remaining fields':'Identity/contact verification; personal time; independent mandatory reflections. No questions were asked; no personal experience was invented.'}
    (SUB/'DELIVERABLES.md').write_text('# Deliverable map\n\n'+'\n'.join('- **'+k+'**: '+v for k,v in required.items())+'\n',encoding='utf-8')
    excluded={'.venv','.tools','.runtime','.git','__pycache__'}
    archive=SUB/'gen-ai-assignment-submission.zip'
    files=[p for p in ROOT.rglob('*') if p.is_file() and not any(part in excluded for part in p.relative_to(ROOT).parts)
           and p.suffix not in ('.zip','.pyc','.db') and not p.name.endswith(('.db-wal','.db-shm')) and p.name!='.env']
    hashes={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in files if p.name!='file_manifest.json'}
    (SUB/'file_manifest.json').write_text(pretty(hashes),encoding='utf-8')
    files=[p for p in files if p.name!='file_manifest.json']+[SUB/'file_manifest.json']
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as bundle:
        for path in files: bundle.write(path,Path('gen-ai-assignment')/path.relative_to(ROOT))
    print('Created '+str(SUB/'completed_submission.docx'),flush=True)
    print(f'Created {archive} ({archive.stat().st_size/1024/1024:.2f} MiB)',flush=True)

if __name__=='__main__': main()
