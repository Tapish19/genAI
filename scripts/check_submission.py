"""Verify the archive, source hashes, original template and actual audit references."""
import hashlib
import json
from pathlib import Path
import re
import zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[1]

def main():
    archive=ROOT/'submission/gen-ai-assignment-submission.zip'
    manifest=json.loads((ROOT/'submission/file_manifest.json').read_text(encoding='utf-8'))
    with zipfile.ZipFile(archive) as bundle:
        assert bundle.testzip() is None,'ZIP checksum failure'
        names=bundle.namelist()
        assert len(names)==len(set(names)),'Duplicate ZIP entries'
        for name in names:
            parts=Path(name).parts
            assert not any(p in ('.venv','.runtime','.tools','.git','__pycache__') for p in parts),name
            assert Path(name).name!='.env',name
            assert not name.endswith(('.gguf','.db','.db-wal','.db-shm')),name
        for path,expected in manifest.items():
            data=bundle.read('gen-ai-assignment/'+path)
            assert hashlib.sha256(data).hexdigest()==expected,path
            assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==expected,'Local file changed after packaging: '+path
        assert 'gen-ai-assignment/submission/project-history.bundle' in names
    source=Document(ROOT/'submission/source/submission_template.docx')
    completed=Document(ROOT/'submission/completed_submission.docx')
    headings=lambda doc:[p.text for p in doc.paragraphs if re.match(r'^[12]\.\d',p.text)]
    assert headings(source)==headings(completed),'Template section headers changed'
    assert len(completed.inline_shapes)==6,'Architecture/evidence images missing'
    all_text='\n'.join([p.text for p in completed.paragraphs]+[c.text for table in completed.tables for row in table.rows for c in row.cells])
    assert not re.search(r'\[\s*(?:\.\.\.|your answer|paste|replace|name)',all_text,re.I),'Template placeholder remained'
    audit=[json.loads(line) for line in (ROOT/'artifacts/llm_audit.jsonl').read_text(encoding='utf-8').splitlines()]
    identifiers={r['request_id'] for r in audit}
    linked=[]
    def visit(value):
        if isinstance(value,dict):
            for key,item in value.items():
                if key=='request_ids':
                    for identifier in item:
                        assert identifier in identifiers,identifier
                        linked.append(identifier)
                elif key=='model_request_id':
                    assert item in identifiers,item
                    linked.append(item)
                else: visit(item)
        elif isinstance(value,list):
            for item in value: visit(item)
    reports={name:json.loads((ROOT/'artifacts'/name).read_text(encoding='utf-8')) for name in ['rag_results.json','judge_validation.json','ab_comparison.json','http_smoke.json']}
    for report in reports.values(): visit(report)
    rag=reports['rag_results.json']; ab=reports['ab_comparison.json']; val=reports['judge_validation.json']
    assert len(rag['cases'])==20 and sum(c['answerable'] for c in rag['cases'])==16
    assert len(ab['cases'])==8 and len(val['cases'])==8
    assert reports['http_smoke.json']['status']=='passed'
    if all(c['mitigated']['winner']=='abstain' for c in ab['cases']):
        assert ab['winner']=='inconclusive','All disagreements cannot declare a winner or tie'
    result={'status':'passed','archive_bytes':archive.stat().st_size,'manifest_files_verified':len(manifest),
            'template_headers_preserved':len(headings(source)),'embedded_figures':len(completed.inline_shapes),
            'audit_references_verified':len(set(linked)),'qa_cases':20,'ab_pairs':8,'validation_probes':8,
            'unit_tests':reports['http_smoke.json']['unit_tests'],'judge_comparison':ab['winner'],
            'gold_reference_comparison':ab['gold_reference_comparison'],
            'candidate_review':'Personal details, disclosures and independent-reflection compliance are the candidate’s responsibility.'}
    print(json.dumps(result,indent=2))

if __name__=='__main__': main()
